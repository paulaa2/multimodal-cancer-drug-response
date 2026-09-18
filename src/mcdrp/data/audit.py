
from __future__ import annotations

import argparse
import csv
import json
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd


@dataclass(frozen=True)
class ColumnRoleAudit:
    """Detected source column for a semantic role."""

    role: str
    detected_column: str | None
    candidates: list[str]


@dataclass(frozen=True)
class TableAudit:
    """Audit information for one configured source table."""

    name: str
    path: str
    description: str
    exists: bool
    n_rows: int | None
    n_columns: int | None
    columns: list[str]
    column_roles: list[ColumnRoleAudit]
    error: str | None = None


@dataclass(frozen=True)
class TablePreview:
    """Lightweight table information used by the audit."""

    n_rows: int
    n_columns: int
    columns: list[str]


def load_audit_config(path: str | Path) -> dict[str, Any]:
    """Load a JSON audit configuration.

    JSON is used here so the very first audit can run with only the standard
    library plus pandas. The rest of the project can still use YAML configs once
    the environment is installed.
    """

    config_path = Path(path)
    with config_path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def read_table(path: Path) -> pd.DataFrame:
    """Read a CSV, TSV, Excel, or parquet table."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix in {".tsv", ".txt"}:
        return pd.read_csv(path, sep="\t")
    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    if suffix in {".parquet", ".pq"}:
        return pd.read_parquet(path)

    raise ValueError(f"Unsupported table format: {path.suffix}")


def read_selected_columns(path: Path, columns: list[str]) -> pd.DataFrame:
    """Read only selected columns from a supported table."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, usecols=columns)
    if suffix in {".tsv", ".txt"}:
        return pd.read_csv(path, sep="\t", usecols=columns)
    if suffix in {".xlsx", ".xls"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        index_by_name = {str(value): index for index, value in enumerate(header)}
        missing = [column for column in columns if column not in index_by_name]
        if missing:
            workbook.close()
            raise ValueError(f"Missing columns in {path}: {missing}")

        selected_indices = [index_by_name[column] for column in columns]
        rows = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            rows.append(
                {
                    column: values[index] if index < len(values) else None
                    for column, index in zip(columns, selected_indices, strict=True)
                }
            )
        workbook.close()
        return pd.DataFrame(rows, columns=columns)
    if suffix in {".parquet", ".pq"}:
        return pd.read_parquet(path, columns=columns)

    raise ValueError(f"Unsupported table format: {path.suffix}")


def inspect_table(path: Path) -> TablePreview:
    """Inspect shape and columns without loading large tables unnecessarily."""

    suffix = path.suffix.lower()
    if suffix == ".csv":
        header = pd.read_csv(path, nrows=0)
        n_rows = count_text_data_rows(path)
        return TablePreview(
            n_rows=n_rows,
            n_columns=len(header.columns),
            columns=list(header.columns),
        )
    if suffix in {".tsv", ".txt"}:
        header = pd.read_csv(path, sep="\t", nrows=0)
        n_rows = count_text_data_rows(path)
        return TablePreview(
            n_rows=n_rows,
            n_columns=len(header.columns),
            columns=list(header.columns),
        )
    if suffix in {".xlsx", ".xls"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = [str(value) for value in first_row]
        n_rows = max((worksheet.max_row or 1) - 1, 0)
        workbook.close()
        return TablePreview(
            n_rows=n_rows,
            n_columns=len(columns),
            columns=columns,
        )
    if suffix in {".parquet", ".pq"}:
        table = pd.read_parquet(path)
        return TablePreview(
            n_rows=len(table),
            n_columns=len(table.columns),
            columns=list(table.columns),
        )

    raise ValueError(f"Unsupported table format: {path.suffix}")


def count_text_data_rows(path: Path) -> int:
    """Count data rows in a delimited text file, excluding the header."""

    with path.open("rb") as handle:
        line_count = sum(1 for _ in handle)
    return max(line_count - 1, 0)


def detect_column(columns: list[str], candidates: list[str]) -> str | None:
    """Find the first candidate column present in a table."""

    exact = {column: column for column in columns}
    lower = {column.lower(): column for column in columns}

    for candidate in candidates:
        if candidate in exact:
            return exact[candidate]
        if candidate.lower() in lower:
            return lower[candidate.lower()]

    return None


def audit_source(name: str, source: dict[str, Any], root: Path) -> TableAudit:
    """Audit one configured source table."""

    relative_path = source["path"]
    path = root / relative_path
    description = source.get("description", "")
    configured_roles = source.get("column_roles", {})

    if not path.exists():
        roles = [
            ColumnRoleAudit(role=role, detected_column=None, candidates=candidates)
            for role, candidates in configured_roles.items()
        ]
        return TableAudit(
            name=name,
            path=relative_path,
            description=description,
            exists=False,
            n_rows=None,
            n_columns=None,
            columns=[],
            column_roles=roles,
        )

    try:
        preview = inspect_table(path)
    except Exception as exc:  # pragma: no cover - message is report output
        return TableAudit(
            name=name,
            path=relative_path,
            description=description,
            exists=True,
            n_rows=None,
            n_columns=None,
            columns=[],
            column_roles=[],
            error=str(exc),
        )

    columns = preview.columns
    roles = [
        ColumnRoleAudit(
            role=role,
            detected_column=detect_column(columns, candidates),
            candidates=candidates,
        )
        for role, candidates in configured_roles.items()
    ]

    return TableAudit(
        name=name,
        path=relative_path,
        description=description,
        exists=True,
        n_rows=preview.n_rows,
        n_columns=preview.n_columns,
        columns=columns,
        column_roles=roles,
    )


def role_map(audit: TableAudit) -> dict[str, str]:
    """Return detected role-to-column mappings for one audited table."""

    return {
        role.role: role.detected_column
        for role in audit.column_roles
        if role.detected_column is not None
    }


def write_cell_line_mapping_template(
    audits: dict[str, TableAudit],
    config: dict[str, Any],
    root: Path,
) -> None:
    response_audit = audits.get("gdsc_response")
    metadata_audit = audits.get("depmap_metadata")
    output_path = root / "data/mappings/cell_line_mapping_template.csv"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, str]] = []

    if response_audit and response_audit.exists and response_audit.error is None:
        source = config["sources"]["gdsc_response"]
        roles = role_map(response_audit)
        id_col = roles.get("cell_line_id")
        name_col = roles.get("cell_line_name")
        if id_col or name_col:
            subset_cols = [col for col in [id_col, name_col] if col]
            table = read_selected_columns(root / source["path"], subset_cols)
            for item in table[subset_cols].drop_duplicates().to_dict("records"):
                rows.append(
                    {
                        "source": "gdsc_response",
                        "original_cell_line_id": str(item.get(id_col, "")),
                        "original_cell_line_name": str(item.get(name_col, "")),
                        "sanger_model_id": str(item.get(id_col, "")),
                        "depmap_id": "",
                        "mapping_status": "unreviewed",
                        "notes": "",
                    }
                )

    if metadata_audit and metadata_audit.exists and metadata_audit.error is None:
        source = config["sources"]["depmap_metadata"]
        roles = role_map(metadata_audit)
        depmap_col = roles.get("depmap_id")
        sanger_col = roles.get("sanger_model_id")
        name_col = roles.get("cell_line_name")
        if depmap_col or sanger_col or name_col:
            subset_cols = [col for col in [depmap_col, sanger_col, name_col] if col]
            table = read_selected_columns(root / source["path"], subset_cols)
            for item in table[subset_cols].drop_duplicates().to_dict("records"):
                rows.append(
                    {
                        "source": "depmap_metadata",
                        "original_cell_line_id": str(item.get(depmap_col, "")),
                        "original_cell_line_name": str(item.get(name_col, "")),
                        "sanger_model_id": str(item.get(sanger_col, "")),
                        "depmap_id": str(item.get(depmap_col, "")),
                        "mapping_status": "reference",
                        "notes": "",
                    }
                )

    write_csv(
        output_path,
        rows,
        fieldnames=[
            "source",
            "original_cell_line_id",
            "original_cell_line_name",
            "sanger_model_id",
            "depmap_id",
            "mapping_status",
            "notes",
        ],
    )


def write_drug_mapping_template(
    audits: dict[str, TableAudit],
    config: dict[str, Any],
    root: Path,
) -> None:
    response_audit = audits.get("gdsc_response")
    structure_audit = audits.get("drug_structures")
    output_path = root / "data/mappings/drug_mapping_template.csv"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, str]] = []

    if response_audit and response_audit.exists and response_audit.error is None:
        source = config["sources"]["gdsc_response"]
        roles = role_map(response_audit)
        id_col = roles.get("drug_id")
        name_col = roles.get("drug_name")
        if id_col or name_col:
            subset_cols = [col for col in [id_col, name_col] if col]
            table = read_selected_columns(root / source["path"], subset_cols)
            for item in table[subset_cols].drop_duplicates().to_dict("records"):
                rows.append(
                    {
                        "source": "gdsc_response",
                        "original_drug_id": str(item.get(id_col, "")),
                        "original_drug_name": str(item.get(name_col, "")),
                        "canonical_smiles": "",
                        "mapping_status": "unreviewed",
                        "notes": "",
                    }
                )

    if structure_audit and structure_audit.exists and structure_audit.error is None:
        source = config["sources"]["drug_structures"]
        roles = role_map(structure_audit)
        id_col = roles.get("drug_id")
        name_col = roles.get("drug_name")
        smiles_col = roles.get("smiles")
        if id_col or name_col or smiles_col:
            subset_cols = [col for col in [id_col, name_col, smiles_col] if col]
            table = read_selected_columns(root / source["path"], subset_cols)
            for item in table[subset_cols].drop_duplicates().to_dict("records"):
                rows.append(
                    {
                        "source": "drug_structures",
                        "original_drug_id": str(item.get(id_col, "")),
                        "original_drug_name": str(item.get(name_col, "")),
                        "canonical_smiles": str(item.get(smiles_col, "")),
                        "mapping_status": "reference",
                        "notes": "",
                    }
                )

    write_csv(
        output_path,
        rows,
        fieldnames=[
            "source",
            "original_drug_id",
            "original_drug_name",
            "canonical_smiles",
            "mapping_status",
            "notes",
        ],
    )


def write_csv(
    path: Path,
    rows: list[dict[str, str]],
    fieldnames: list[str],
) -> None:

    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def run_audit(config_path: str | Path, root: str | Path = ".") -> dict[str, Any]:
    """Run the configured data audit and write report files."""

    root_path = Path(root).resolve()
    config = load_audit_config(config_path)

    audits = {
        name: audit_source(name, source, root_path)
        for name, source in config.get("sources", {}).items()
    }

    report = {
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "project": config.get("project"),
        "sources": {name: asdict(audit) for name, audit in audits.items()},
    }

    manifest_dir = root_path / "data/manifests"
    manifest_dir.mkdir(parents=True, exist_ok=True)
    report_path = manifest_dir / "data_audit_report.json"
    with report_path.open("w", encoding="utf-8") as handle:
        json.dump(report, handle, indent=2)

    write_cell_line_mapping_template(audits, config, root_path)
    write_drug_mapping_template(audits, config, root_path)

    return report


def build_parser() -> argparse.ArgumentParser:
    """Create the command-line parser."""

    parser = argparse.ArgumentParser(description="Run the F0 data audit.")
    parser.add_argument(
        "--config",
        default="configs/data_audit.example.json",
        help="Path to the data-audit JSON config.",
    )
    parser.add_argument(
        "--root",
        default=".",
        help="Project root. Defaults to the current directory.",
    )
    return parser


def main() -> None:
    """Command-line entry point."""

    args = build_parser().parse_args()
    report = run_audit(args.config, args.root)

    print("Data audit complete")
    for name, source in report["sources"].items():
        status = "found" if source["exists"] else "missing"
        shape = ""
        if source["n_rows"] is not None:
            shape = f" ({source['n_rows']} rows x {source['n_columns']} columns)"
        print(f"- {name}: {status}{shape}")


if __name__ == "__main__":
    main()
